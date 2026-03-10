"""XLSX extraction — openpyxl for all sheets."""

import io
from typing import Any

from openpyxl import load_workbook


async def extract_xlsx(file_bytes: bytes, filename: str) -> dict[str, Any]:
    """Extract all sheets from an XLSX file as text."""
    wb = load_workbook(io.BytesIO(file_bytes), read_only=True, data_only=True)
    parts: list[str] = []

    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        parts.append(f"## Sheet: {sheet_name}")
        rows: list[str] = []
        for row in ws.iter_rows(values_only=True):
            cells = [str(c) if c is not None else "" for c in row]
            # Skip completely empty rows
            if any(c.strip() for c in cells):
                rows.append("| " + " | ".join(cells) + " |")
        if rows:
            # Add markdown table header separator after first row
            first_row_cols = rows[0].count("|") - 1
            header_sep = "| " + " | ".join(["---"] * max(first_row_cols, 1)) + " |"
            rows.insert(1, header_sep)
            parts.append("\n".join(rows))

    wb.close()
    text = "\n\n".join(parts)

    return {
        "filename": filename,
        "content_type": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        "text": text if text else None,
        "images": [],
        "pages": None,
        "characters": len(text),
        "method": "openpyxl",
    }
