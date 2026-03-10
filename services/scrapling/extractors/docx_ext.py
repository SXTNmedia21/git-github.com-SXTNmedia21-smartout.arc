"""DOCX extraction — python-docx for paragraphs and tables."""

import io
from typing import Any

from docx import Document


async def extract_docx(file_bytes: bytes, filename: str) -> dict[str, Any]:
    """Extract text and tables from a DOCX file."""
    doc = Document(io.BytesIO(file_bytes))
    parts: list[str] = []

    for para in doc.paragraphs:
        text = para.text.strip()
        if not text:
            continue
        # Map heading styles to markdown
        if para.style and para.style.name.startswith("Heading"):
            level = para.style.name.replace("Heading ", "").strip()
            try:
                hashes = "#" * int(level)
            except ValueError:
                hashes = "##"
            parts.append(f"{hashes} {text}")
        else:
            parts.append(text)

    # Tables as markdown
    for table in doc.tables:
        rows: list[str] = []
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            rows.append("| " + " | ".join(cells) + " |")
        if rows:
            col_count = len(table.rows[0].cells)
            header_sep = "| " + " | ".join(["---"] * col_count) + " |"
            rows.insert(1, header_sep)
            parts.append("\n".join(rows))

    text = "\n\n".join(parts)

    return {
        "filename": filename,
        "content_type": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        "text": text if text else None,
        "images": [],
        "pages": None,
        "characters": len(text),
        "method": "python-docx",
    }
